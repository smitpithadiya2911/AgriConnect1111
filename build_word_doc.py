import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_report():
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    PRIMARY_COLOR = RGBColor(27, 67, 50)      # Deep Forest Green (#1B4332)
    SECONDARY_COLOR = RGBColor(45, 106, 79)  # Medium Green (#2D6A4F)
    TEXT_DARK = RGBColor(33, 33, 33)         # Dark Charcoal (#212121)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
        tblPr = table._element.xpath('w:tblPr')
        if tblPr:
            borders = parse_xml(f'''
                <w:tblBorders {nsdecls("w")}>
                    <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                    <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                    <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                    <w:insideV w:val="none"/>
                    <w:left w:val="none"/>
                    <w:right w:val="none"/>
                </w:tblBorders>
            ''')
            tblPr[0].append(borders)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = SECONDARY_COLOR
        return p

    def add_paragraph(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
            r_pre.font.color.rgb = TEXT_DARK
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_DARK
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
            r_pre.font.color.rgb = TEXT_DARK
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_DARK
        return p

    # -------------------------------------------------------------
    # PAGE 1: TITLE PAGE
    # -------------------------------------------------------------
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(36)
    r = p_top.add_run("A Project Report On…")
    r.font.name = 'Calibri'
    r.font.size = Pt(18)
    r.font.italic = True
    r.font.color.rgb = SECONDARY_COLOR

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(36)
    r_title = p_title.add_run("AgriConnect: Digital Agricultural Marketplace & Direct Farmer-to-Buyer Platform")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    p_dev = doc.add_paragraph()
    p_dev.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dev.paragraph_format.space_after = Pt(36)
    r_dev = p_dev.add_run("Developed By…\nSmit Pithadiya\n(BCA Semester - 6)")
    r_dev.font.name = 'Calibri'
    r_dev.font.size = Pt(16)
    r_dev.font.bold = True
    r_dev.font.color.rgb = TEXT_DARK

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Submitted To..\nGeetanjali College Of Computer Science And Commerce ( B.C.A. )\nRajkot\nSaurashtra University Rajkot\nAcademic Year: 2025 - 2026")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = TEXT_DARK

    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_guide.paragraph_format.space_before = Pt(36)
    r_guide = p_guide.add_run("Project Guide : PROF. Harsh Joshi")
    r_guide.font.name = 'Calibri'
    r_guide.font.size = Pt(14)
    r_guide.font.bold = True
    r_guide.font.color.rgb = PRIMARY_COLOR

    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 2: ACKNOWLEDGEMENTS
    # -------------------------------------------------------------
    add_heading_1("Acknowledgements")
    add_bullet("I am happy to submit my idea and implementation for the \"AgriConnect - Agricultural Marketplace System\" web application to Saurashtra University, Rajkot for a BCA degree in the Computer Science branch.")
    add_bullet("I am also deeply grateful to Prof. Brijesh Shah, Head of the Department, and all faculty members of the Department of Computer Science for their kind guidance, constant encouragement, and technical support throughout this journey.")
    add_bullet("I take the privilege to acknowledge the authors of numerous reference books, research publications, Django official documentation, and technical web blogs which we referred to during the development of this project.")
    add_bullet("I express my heartfelt gratitude to my parents and family for their unwavering moral and financial support, without which I could accomplish nothing in my academic pursuits and personal life.")
    add_bullet("The feeling of gratefulness for everyone's timely assistance directly arises from the bottom of my heart. Small but critical guidance during testing and database schema design proved to be significant milestones in completing this application.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 3: PREFACE
    # -------------------------------------------------------------
    add_heading_1("Preface")
    add_bullet("Practical training and software implementation are of paramount importance in understanding theoretical aspects of Computer Application & Web Software Development. Recognizing this importance, we prepared the \"AgriConnect\" project to enrich our practical knowledge regarding web technologies, full-stack architecture, and database management systems.")
    add_bullet("We are very much pleased to present this comprehensive project report on \"AgriConnect: Digital Agricultural Marketplace\", developed at Geetanjali College of Computer Science and Commerce, affiliated with Saurashtra University.")
    add_bullet("This report contains an exhaustive technical and functional overview of the entire project. Anyone with basic computer software knowledge can easily understand the contents of this report. System design is illustrated using Data Flow Diagrams (DFDs), Entity Relationship (ER) diagrams, exact Data Dictionary schema definitions, test cases, and screen illustrations.")
    add_bullet("We have tried our best in practical study, systematic code structuring, and clear presentation of this report.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 4: INDEX / TABLE OF CONTENTS
    # -------------------------------------------------------------
    add_heading_1("Index")
    
    table_idx = doc.add_table(rows=1, cols=2)
    table_idx.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_idx)
    
    hdr_cells = table_idx.rows[0].cells
    hdr_cells[0].text = "Topic / Section Title"
    hdr_cells[1].text = "Page No."
    set_cell_background(hdr_cells[0], "1B4332")
    set_cell_background(hdr_cells[1], "1B4332")
    for cell in hdr_cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    index_items = [
        ("1. Project Profile", "5"),
        ("2. Software Development Life Cycle (SDLC)", "6"),
        ("3. System Requirements (Hardware & Software)", "8"),
        ("4. Development & Coding Standards", "9"),
        ("5. About the Tool (Tools & Technologies Used)", "11"),
        ("6. Data Flow Diagrams (DFD Level 0 & Level 1)", "13"),
        ("7. Entity Relationship (ER) Diagram & Architecture", "15"),
        ("8. Data Dictionary (Complete Database Schema)", "17"),
        ("9. Screenshots & Module Descriptions", "22"),
        ("10. Test Cases (Farmer, Buyer & Admin Modules)", "32"),
        ("11. Limitations", "37"),
        ("12. Future Enhancements", "38"),
        ("13. Webliography", "39")
    ]

    for item, page in index_items:
        row_cells = table_idx.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = page
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 1: PROJECT PROFILE
    # -------------------------------------------------------------
    add_heading_1("1. Project Profile")

    table_prof = doc.add_table(rows=0, cols=2)
    table_prof.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_prof)

    profile_data = [
        ("Project Title:", "AgriConnect: Digital Agricultural Marketplace System"),
        ("Development Software:", "Visual Studio Code, Python 3.12+, Django 5.x Framework"),
        ("Front End:", "HTML5, CSS3, JavaScript (ES6+), Bootstrap 5, FontAwesome"),
        ("Backend / Logic:", "Python, Django Web Framework (MVT Architecture)"),
        ("Database:", "MySQL Database Management System"),
        ("Academic Year:", "2025 - 2026"),
        ("Developed By:", "Smit Pithadiya (BCA Semester - 6)"),
        ("Submitted To:", "Geetanjali College Of Computer Science And Commerce, Rajkot"),
        ("Documentation Tool:", "Microsoft Word, Draw.io, Python Docx, Matplotlib"),
        ("Operating System:", "Cross-platform (Windows 10/11, macOS, Linux)"),
        ("Programming Language:", "Python 3.x, JavaScript, HTML, CSS")
    ]

    for label, val in profile_data:
        row_cells = table_prof.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = val
        set_cell_background(row_cells[0], "F4F9F4")
        row_cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 2: SOFTWARE DEVELOPMENT LIFE CYCLE (SDLC)
    # -------------------------------------------------------------
    add_heading_1("2. Software Development Life Cycle (SDLC)")
    add_paragraph("The Software Development Life Cycle (SDLC) is a systematic, structured methodology utilized to design, develop, test, and deploy high-quality web software applications. For the development of the AgriConnect project, the Waterfall Model was selected due to its linear, sequential flow and well-defined phase dependencies.")
    
    add_heading_2("Waterfall Model Steps:")
    add_bullet("Requirement Gathering & Feasibility Analysis", "1. ")
    add_bullet("Project Planning & Scope Definition", "2. ")
    add_bullet("System Design & Architecture Specification", "3. ")
    add_bullet("Coding & Full-Stack Implementation", "4. ")
    add_bullet("Testing & Quality Assurance", "5. ")

    add_heading_2("Phase 1: Requirement Gathering")
    add_paragraph("In this initial stage, detailed functional and non-functional requirements were gathered by researching existing agricultural trade platforms and consulting local farmers and produce buyers.")
    add_bullet("Farmer Requirements: Registration, identity verification, crop listing management, quantity tracking, price per kg setup, stock status, direct buyer order reception, earnings summary, and direct chat.", "• ")
    add_bullet("Buyer Requirements: Marketplace browsing, category filtering, search, crop details, shopping cart management, delivery address entry, multi-mode payment, order placement with OTP verification, order history, and farmer reviews.", "• ")
    add_bullet("Admin Requirements: Centralized dashboard, crop approval workflow, user account role management, platform analytics, system reports generation, and dispute resolution.", "• ")

    add_heading_2("Phase 2: Project Planning")
    add_paragraph("A comprehensive project plan was established outlining timelines, milestone deadlines, module division, technology selection (Python + Django MVT), and database choice.")

    add_heading_2("Phase 3: System Design")
    add_paragraph("The architecture was designed using Django's Model-View-Template (MVT) pattern. Database schemas were constructed for 21 core entities including User, Farmer, Buyer, AuthOTP, Category, Crop, CartItem, Order, OrderItem, Payment, Review, Notification, ChatMessage, Wishlist, FarmerRating, Feedback, PriceTrend, MarketInsight, ReturnRequest, Report, and OrderOTP tables. Data Flow Diagrams (DFDs) and ER diagrams were modeled to guarantee clean relationships.")

    add_heading_2("Phase 4: Coding & Implementation")
    add_paragraph("Development was executed using Python 3 and Django 5. Web templates were built using HTML5, CSS3, JavaScript, and Bootstrap 5 to deliver a responsive user interface. Clean coding standards (PEP 8) and secure design principles (CSRF tokens, password hashing, RBAC) were strictly enforced.")

    add_heading_2("Phase 5: Testing")
    add_paragraph("Rigorous Unit Testing, Integration Testing, System Testing, and User Acceptance Testing (UAT) were conducted across all farmer, buyer, and admin workflows to verify functional correctness and zero defects.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 3: SYSTEM REQUIREMENTS
    # -------------------------------------------------------------
    add_heading_1("3. System Requirements")
    add_paragraph("To successfully deploy, run, and maintain the AgriConnect project, the following hardware and software specifications are required:")

    add_heading_2("Minimum Hardware Requirements:")
    table_hw = doc.add_table(rows=1, cols=2)
    table_hw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_hw)
    
    r0 = table_hw.rows[0].cells
    r0[0].text = "Component / Resource"
    r0[1].text = "Minimum Specification"
    set_cell_background(r0[0], "1B4332")
    set_cell_background(r0[1], "1B4332")
    for cell in r0:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    hw_specs = [
        ("Operating System", "Windows 10/11, macOS 11+, Linux Ubuntu 20.04+"),
        ("CPU / Processor", "Intel Core i3 / AMD Ryzen 3 Dual-Core (2.0 GHz or higher)"),
        ("System RAM", "4 GB (8 GB recommended for optimal development)"),
        ("Hard Disk Space", "20 GB available storage space"),
        ("Display Resolution", "1280 x 720 HD Resolution or higher"),
        ("Network Connection", "Broadband Internet (for static CDN resources & deployment)")
    ]
    for k, v in hw_specs:
        row_cells = table_hw.add_row().cells
        row_cells[0].text = k
        row_cells[1].text = v
        set_cell_background(row_cells[0], "F4F9F4")

    add_heading_2("Software Requirements:")
    table_sw = doc.add_table(rows=1, cols=2)
    table_sw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_sw)
    
    r0_sw = table_sw.rows[0].cells
    r0_sw[0].text = "Software Tool / Component"
    r0_sw[1].text = "Version / Specification"
    set_cell_background(r0_sw[0], "1B4332")
    set_cell_background(r0_sw[1], "1B4332")
    for cell in r0_sw:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    sw_specs = [
        ("Runtime Environment", "Python 3.12+ 64-bit Interpreter"),
        ("Web Framework", "Django 5.x Framework"),
        ("Database Server", "MySQL Database Server (v8.0+ / TiDB Cloud)"),
        ("Frontend Stack", "HTML5, CSS3, JavaScript (ES6+), Bootstrap 5.3"),
        ("Web Server", "Django WSGI / ASGI Dev Server (Apache 2.4 / Nginx in Production)"),
        ("IDE / Source Editor", "Visual Studio Code / PyCharm"),
        ("Version Control", "Git & GitHub Repository")
    ]
    for k, v in sw_specs:
        row_cells = table_sw.add_row().cells
        row_cells[0].text = k
        row_cells[1].text = v
        set_cell_background(row_cells[0], "F4F9F4")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 4: DEVELOPMENT & CODING STANDARDS
    # -------------------------------------------------------------
    add_heading_1("4. Development & Coding Standards")
    add_paragraph("To ensure maximum readability, maintainability, and scalability across the AgriConnect codebase, strict Python (PEP 8) and web engineering standards were enforced during implementation.")

    add_heading_2("Naming Conventions:")
    add_bullet("Use snake_case for filenames, module names, Python variables, and function names (e.g., crop_detail_view, price_per_kg).", "• File & Function Naming: ")
    add_bullet("Use PascalCase for Django models, views, forms, and Python classes (e.g., Crop, CartItem, Wishlist).", "• Class Naming: ")
    add_bullet("Use lowercase snake_case for HTML template files placed under clear app folders (e.g., templates/marketplace/crop_detail.html).", "• Template Naming: ")

    add_heading_2("Code Formatting & Structure:")
    add_bullet("Consistently use 4 spaces per indentation level. Avoid tab characters to maintain environment consistency.", "• Indentation: ")
    add_bullet("Limit Python source code lines to 79-88 characters for enhanced visual readability and side-by-side editing.", "• Line Length: ")
    add_bullet("Group imports logically: Standard library imports first, Third-party imports second, and Local app imports third.", "• Import Organization: ")

    add_heading_2("Error Handling & Security Standards:")
    add_bullet("Wrap database queries, file operations, and external API requests inside try-except blocks to gracefully capture runtime exceptions.", "• Exception Handling: ")
    add_bullet("Enforce @login_required and role-verification decorators on views to restrict unauthorized URL access.", "• Access Control: ")
    add_bullet("Include {% csrf_token %} on all HTML POST forms to prevent Cross-Site Request Forgery attacks.", "• Security Tokens: ")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 5: ABOUT THE TOOLS & TECHNOLOGIES USED
    # -------------------------------------------------------------
    add_heading_1("5. About the Tool (Tools & Technologies Used)")

    add_heading_2("Python Programming Language:")
    add_paragraph("Python is a high-level, interpreted, object-oriented programming language known for its simple syntax, readability, and extensive standard library. Python powers the backend business logic and data manipulation of AgriConnect.")

    add_heading_2("Django Web Framework:")
    add_paragraph("Django is a high-level Python web framework that encourages rapid development and clean, pragmatic design. Built by experienced developers, it takes care of much of the hassle of web development. Django follows the Model-View-Template (MVT) architectural pattern and provides built-in Object-Relational Mapping (ORM), robust user authentication, form validation, and admin portal capabilities out of the box.")

    add_heading_2("HTML5 & CSS3 & Bootstrap 5:")
    add_paragraph("HTML5 provides semantic structure for web pages, while CSS3 handles responsive visual styling. Bootstrap 5, a popular front-end CSS framework, is integrated to ensure AgriConnect renders flawlessly on desktop displays, tablets, and smartphones.")

    add_heading_2("MySQL Database Management System (DBMS):")
    add_paragraph("MySQL is a high-performance, robust relational database management system (RDBMS) used to store and manage all relational data in AgriConnect. Structured database tables store user credentials, farmer profiles, crop inventories, shopping cart items, order transactions, payments, reviews, and market insights. In production, MySQL / TiDB Cloud is utilized to deliver enterprise-grade performance, ACID transaction safety, foreign key integrity, and high-concurrency reliability.")

    add_heading_2("Visual Studio Code (VS Code):")
    add_paragraph("VS Code is a lightweight but powerful source code editor running on desktop OS. It features built-in support for Python, HTML/CSS, Git version control, syntax highlighting, and interactive debugging.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 6: DATA FLOW DIAGRAMS (DFD)
    # -------------------------------------------------------------
    add_heading_1("6. Data Flow Diagrams (DFD)")
    add_paragraph("Data Flow Diagrams graphically illustrate how data flows through the AgriConnect system, identifying external entities, process transformations, and internal database stores.")

    add_heading_2("Data Flow Diagram (DFD Level 0 - Context Diagram):")
    if os.path.exists('report_assets/dfd_level0.png'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture('report_assets/dfd_level0.png', width=Inches(5.8))
    add_paragraph("Figure 6.1: Context Diagram (DFD Level 0) illustrating high-level interactions between Farmers, Buyers, Admins, and the central AgriConnect platform.")

    doc.add_page_break()

    add_heading_2("Data Flow Diagram (DFD Level 1 - Process Breakdown):")
    if os.path.exists('report_assets/dfd_level1.png'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture('report_assets/dfd_level1.png', width=Inches(5.8))
    add_paragraph("Figure 6.2: Detailed Process Breakdown (DFD Level 1) showing data paths between authentication, crop listing, shopping cart, order placement with OTP, payment, and database stores.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 7: ENTITY RELATIONSHIP (ER) DIAGRAM & ARCHITECTURE
    # -------------------------------------------------------------
    add_heading_1("7. Entity Relationship (ER) Diagram & Architecture")
    
    add_heading_2("Entity Relationship (ER) Diagram:")
    if os.path.exists('report_assets/er_diagram.png'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture('report_assets/er_diagram.png', width=Inches(5.8))
    add_paragraph("Figure 7.1: Entity Relationship Diagram depicting primary keys, foreign key constraints, and cardinalities between User, Crop, Category, Order, OrderItem, and Payment entities.")

    doc.add_page_break()

    add_heading_2("System Architecture Diagram:")
    if os.path.exists('report_assets/architecture_diagram.png'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture('report_assets/architecture_diagram.png', width=Inches(5.8))
    add_paragraph("Figure 7.2: 3-Tier Architecture showing Presentation Layer, Django Application Server Layer, and Relational Database Storage Layer.")

    doc.add_page_break()

    add_heading_2("Use Case Diagram:")
    if os.path.exists('report_assets/use_case_diagram.png'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture('report_assets/use_case_diagram.png', width=Inches(5.8))
    add_paragraph("Figure 7.3: Use Case Diagram depicting interactions between Farmer, Buyer, and Admin roles.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 8: DATA DICTIONARY
    # -------------------------------------------------------------
    add_heading_1("8. Data Dictionary")
    add_paragraph("The Data Dictionary provides the exact schema definition for all 21 database tables implemented in the AgriConnect relational database system, detailing field names, data types, nullability, key constraints, and functional descriptions.")

    models_info = [
        ("Table 1: accounts_user (Custom User Model)", [
            ("id", "BigAutoField", "No", "Primary Key", "Unique user identifier"),
            ("username", "CharField", "No", "Unique", "Login username"),
            ("email", "EmailField", "No", "Unique", "User email address"),
            ("role", "CharField", "No", "-", "User role (admin, farmer, buyer)"),
            ("phone", "CharField", "Yes", "-", "Contact phone number"),
            ("address", "TextField", "Yes", "-", "Physical delivery/farm address"),
            ("profile_picture", "ImageField", "Yes", "-", "Profile avatar image"),
            ("is_active", "BooleanField", "No", "-", "Account active status flag"),
            ("date_joined", "DateTimeField", "No", "-", "Account registration timestamp")
        ]),
        ("Table 2: accounts_farmer (Farmer Profile)", [
            ("id", "BigAutoField", "No", "Primary Key", "Profile ID"),
            ("user_id", "OneToOneField", "No", "FK (User)", "Reference to accounts_user"),
            ("farm_name", "CharField", "Yes", "-", "Name of the farm / enterprise"),
            ("farm_location", "CharField", "Yes", "-", "Farm address / region"),
            ("farm_size", "CharField", "Yes", "-", "Land acreage (e.g. 5 Acres)"),
            ("certification_status", "CharField", "No", "-", "Verification badge status (verified/pending/none)"),
            ("experience", "CharField", "No", "-", "Years of agricultural experience"),
            ("orders_completed", "IntegerField", "No", "-", "Count of fulfilled orders"),
            ("store_rating", "FloatField", "No", "-", "Average store rating (0.0 - 5.0)"),
            ("specialization", "CharField", "No", "-", "Farming domain specialization"),
            ("about", "TextField", "Yes", "-", "Farmer biography & overview"),
            ("cover_banner", "ImageField", "Yes", "-", "Storefront header cover banner"),
            ("certifications", "TextField", "Yes", "-", "Agricultural certifications held"),
            ("awards", "TextField", "Yes", "-", "Recognitions and awards received"),
            ("repeat_customers", "IntegerField", "No", "-", "Count/rate of repeat buyers"),
            ("ontime_delivery", "IntegerField", "No", "-", "On-time delivery percentage"),
            ("satisfaction_rate", "IntegerField", "No", "-", "Customer satisfaction percentage"),
            ("village", "CharField", "Yes", "-", "Village name"),
            ("city", "CharField", "Yes", "-", "City location"),
            ("state", "CharField", "Yes", "-", "State location"),
            ("pincode", "CharField", "Yes", "-", "Postal code"),
            ("organic_certificate", "FileField", "Yes", "-", "Uploaded organic farming certificate"),
            ("verification_status", "CharField", "No", "-", "Admin verification status choice"),
            ("aadhaar_document", "FileField", "Yes", "-", "Identity verification document"),
            ("farm_certificate", "FileField", "Yes", "-", "Land certificate document"),
            ("trust_score", "IntegerField", "No", "-", "Platform trust rating score"),
            ("verification_date", "DateTimeField", "Yes", "-", "Verification approval timestamp"),
            ("verified_by_id", "ForeignKey", "Yes", "FK (User)", "Admin user performing verification"),
            ("admin_notes", "TextField", "Yes", "-", "Verification review notes"),
            ("approval_history", "TextField", "Yes", "-", "Audit trail log of actions")
        ]),
        ("Table 3: accounts_buyer (Buyer Profile)", [
            ("id", "BigAutoField", "No", "Primary Key", "Profile ID"),
            ("user_id", "OneToOneField", "No", "FK (User)", "Reference to accounts_user"),
            ("delivery_address", "TextField", "Yes", "-", "Primary shipping destination"),
            ("contact_name", "CharField", "Yes", "-", "Contact person name"),
            ("city", "CharField", "Yes", "-", "Buyer city location"),
            ("state", "CharField", "Yes", "-", "Buyer state location"),
            ("pincode", "CharField", "Yes", "-", "Postal pincode")
        ]),
        ("Table 4: accounts_authotp (Authentication OTP)", [
            ("id", "BigAutoField", "No", "Primary Key", "OTP Record ID"),
            ("user_id", "ForeignKey", "Yes", "FK (User)", "Target user reference"),
            ("email", "EmailField", "Yes", "-", "Email address for unregistered OTP"),
            ("otp_code", "CharField", "No", "-", "6-digit OTP passcode"),
            ("otp_type", "CharField", "No", "-", "OTP type (registration / password_reset)"),
            ("created_at", "DateTimeField", "No", "-", "OTP generation timestamp"),
            ("expires_at", "DateTimeField", "No", "-", "OTP expiration timestamp"),
            ("is_verified", "BooleanField", "No", "-", "OTP verification status flag"),
            ("attempts", "IntegerField", "No", "-", "Failed attempt counter"),
            ("resend_count", "IntegerField", "No", "-", "Resend request counter")
        ]),
        ("Table 5: marketplace_category (Crop Category)", [
            ("id", "BigAutoField", "No", "Primary Key", "Category ID"),
            ("name", "CharField", "No", "Unique", "Category title (Vegetables, Fruits, Grains)"),
            ("description", "TextField", "Yes", "-", "Category detailed description"),
            ("image", "ImageField", "Yes", "-", "Category banner thumbnail image")
        ]),
        ("Table 6: marketplace_crop (Crop Product Listing)", [
            ("id", "BigAutoField", "No", "Primary Key", "Crop Product ID"),
            ("farmer_id", "ForeignKey", "No", "FK (Farmer)", "Farmer owner of crop listing"),
            ("category_id", "ForeignKey", "No", "FK (Category)", "Produce category reference"),
            ("name", "CharField", "No", "-", "Crop name (e.g. Organic Wheat)"),
            ("description", "TextField", "No", "-", "Detailed produce description"),
            ("price_per_kg", "DecimalField", "No", "-", "Price per unit / kilogram (INR)"),
            ("quantity_available", "DecimalField", "No", "-", "Available stock quantity"),
            ("unit", "CharField", "No", "-", "Unit choice (kg, ton, piece, quintal)"),
            ("image", "ImageField", "Yes", "-", "Produce primary photograph"),
            ("is_approved", "BooleanField", "No", "-", "Admin approval status"),
            ("harvest_date", "DateField", "No", "-", "Date crop was harvested"),
            ("shelf_life_days", "IntegerField", "No", "-", "Estimated shelf life in days"),
            ("availability_status", "CharField", "No", "-", "Stock status (available/out_of_stock)"),
            ("created_at", "DateTimeField", "No", "-", "Listing creation timestamp"),
            ("updated_at", "DateTimeField", "No", "-", "Listing last update timestamp"),
            ("status_history", "JSONField", "No", "-", "Audit trail of listing changes"),
            ("otp_code", "CharField", "Yes", "-", "Verification OTP code"),
            ("otp_expires_at", "DateTimeField", "Yes", "-", "OTP expiry timestamp"),
            ("is_otp_verified", "BooleanField", "No", "-", "OTP verification flag"),
            ("return_reason", "CharField", "Yes", "-", "Return reason choice if returned"),
            ("return_description", "TextField", "Yes", "-", "Return details note"),
            ("return_status", "CharField", "Yes", "-", "Return status (Pending/Approved/Rejected)"),
            ("market_demand_level", "CharField", "No", "-", "Market demand index (High/Medium/Low)"),
            ("market_forecast_direction", "CharField", "No", "-", "Price forecast direction"),
            ("current_market_price", "DecimalField", "Yes", "-", "APMC mandi benchmark price"),
            ("market_price_history", "JSONField", "No", "-", "Price history trend data")
        ]),
        ("Table 7: marketplace_cartitem (Shopping Cart Item)", [
            ("id", "BigAutoField", "No", "Primary Key", "Cart Item ID"),
            ("user_id", "ForeignKey", "No", "FK (User)", "Buyer user owning cart"),
            ("crop_id", "ForeignKey", "No", "FK (Crop)", "Selected crop item"),
            ("quantity", "PositiveInt", "No", "-", "Selected quantity count")
        ]),
        ("Table 8: marketplace_order (Order Master Table)", [
            ("id", "BigAutoField", "No", "Primary Key", "Order ID"),
            ("buyer_id", "ForeignKey", "No", "FK (User)", "Buyer placing the order"),
            ("farmer_id", "ForeignKey", "No", "FK (Farmer)", "Farmer fulfilling order"),
            ("total_amount", "DecimalField", "No", "-", "Total payable order cost (INR)"),
            ("payment_method", "CharField", "No", "-", "Payment method choice (COD/Online)"),
            ("status", "CharField", "No", "-", "Order status (Pending/Confirmed/Delivered)"),
            ("shipping_address", "TextField", "No", "-", "Complete shipping destination"),
            ("invoice_pdf", "FileField", "Yes", "-", "Generated PDF tax invoice"),
            ("created_at", "DateTimeField", "No", "-", "Order placement timestamp"),
            ("updated_at", "DateTimeField", "No", "-", "Order last modified timestamp"),
            ("status_history", "JSONField", "No", "-", "Order state lifecycle history"),
            ("otp_code", "CharField", "Yes", "-", "6-digit delivery confirmation OTP"),
            ("otp_expires_at", "DateTimeField", "Yes", "-", "OTP expiry timestamp"),
            ("is_otp_verified", "BooleanField", "No", "-", "OTP verified status flag"),
            ("return_reason", "CharField", "Yes", "-", "Order return reason"),
            ("return_description", "TextField", "Yes", "-", "Order return details"),
            ("return_status", "CharField", "Yes", "-", "Return request approval status")
        ]),
        ("Table 9: marketplace_orderitem (Order Line Items)", [
            ("id", "BigAutoField", "No", "Primary Key", "Line Item ID"),
            ("order_id", "ForeignKey", "No", "FK (Order)", "Parent order reference"),
            ("crop_id", "ForeignKey", "Yes", "FK (Crop)", "Purchased crop product"),
            ("quantity", "DecimalField", "No", "-", "Purchased quantity volume"),
            ("price_per_unit", "DecimalField", "No", "-", "Unit price at purchase time")
        ]),
        ("Table 10: marketplace_payment (Payment Transaction)", [
            ("id", "BigAutoField", "No", "Primary Key", "Payment ID"),
            ("order_id", "OneToOneField", "No", "FK (Order)", "Associated order reference"),
            ("payment_method", "CharField", "No", "-", "Payment mode (COD / Online / UPI)"),
            ("transaction_id", "CharField", "Yes", "-", "Gateway transaction reference"),
            ("status", "CharField", "No", "-", "Payment status (Pending/Completed/Failed)"),
            ("amount", "DecimalField", "No", "-", "Transaction total amount"),
            ("created_at", "DateTimeField", "No", "-", "Payment processing timestamp")
        ]),
        ("Table 11: marketplace_review (Customer Crop & Farmer Review)", [
            ("id", "BigAutoField", "No", "Primary Key", "Review ID"),
            ("buyer_id", "ForeignKey", "No", "FK (User)", "Reviewer buyer user"),
            ("crop_id", "ForeignKey", "Yes", "FK (Crop)", "Reviewed crop product"),
            ("farmer_id", "ForeignKey", "Yes", "FK (Farmer)", "Reviewed farmer profile"),
            ("quality_rating", "PositiveInt", "Yes", "-", "Crop quality rating (1-5)"),
            ("communication_rating", "PositiveInt", "Yes", "-", "Farmer response score (1-5)"),
            ("delivery_rating", "PositiveInt", "Yes", "-", "Delivery speed score (1-5)"),
            ("packaging_rating", "PositiveInt", "Yes", "-", "Packaging quality score (1-5)"),
            ("rating", "PositiveInt", "No", "-", "Overall star rating score (1-5)"),
            ("title", "CharField", "Yes", "-", "Review headline title"),
            ("comment", "TextField", "Yes", "-", "Detailed feedback comment"),
            ("image", "ImageField", "Yes", "-", "Customer product photo upload"),
            ("reply", "TextField", "Yes", "-", "Seller response reply text"),
            ("reply_at", "DateTimeField", "Yes", "-", "Seller reply timestamp"),
            ("created_at", "DateTimeField", "No", "-", "Review submission timestamp")
        ]),
        ("Table 12: marketplace_notification (System Notification)", [
            ("id", "BigAutoField", "No", "Primary Key", "Notification ID"),
            ("user_id", "ForeignKey", "No", "FK (User)", "Recipient user account"),
            ("message", "TextField", "No", "-", "Notification text body"),
            ("notification_type", "CharField", "No", "-", "Type (order/price/weather/system/chat/review/verification)"),
            ("priority", "CharField", "No", "-", "Priority level (high/medium/low)"),
            ("link", "CharField", "Yes", "-", "Action link URL"),
            ("is_read", "BooleanField", "No", "-", "Read/unread status flag"),
            ("created_at", "DateTimeField", "No", "-", "Dispatch timestamp")
        ]),
        ("Table 13: marketplace_chatmessage (Buyer-Farmer Chat Message)", [
            ("id", "BigAutoField", "No", "Primary Key", "Message ID"),
            ("sender_id", "ForeignKey", "No", "FK (User)", "Sender user reference"),
            ("receiver_id", "ForeignKey", "No", "FK (User)", "Recipient user reference"),
            ("crop_id", "ForeignKey", "Yes", "FK (Crop)", "Associated crop produce inquiry"),
            ("message", "TextField", "No", "-", "Message text content"),
            ("is_read", "BooleanField", "No", "-", "Read status flag"),
            ("created_at", "DateTimeField", "No", "-", "Message sent timestamp")
        ]),
        ("Table 14: marketplace_wishlist (Saved Crops Wishlist)", [
            ("id", "BigAutoField", "No", "Primary Key", "Wishlist Item ID"),
            ("user_id", "ForeignKey", "No", "FK (User)", "Buyer user reference"),
            ("crop_id", "ForeignKey", "No", "FK (Crop)", "Saved crop listing reference"),
            ("added_at", "DateTimeField", "No", "-", "Added timestamp")
        ]),
        ("Table 15: marketplace_farmerrating (Farmer Rating & Evaluation)", [
            ("id", "BigAutoField", "No", "Primary Key", "Rating Record ID"),
            ("buyer_id", "ForeignKey", "No", "FK (User)", "Rating buyer reference"),
            ("farmer_id", "ForeignKey", "No", "FK (Farmer)", "Rated farmer reference"),
            ("rating", "PositiveInt", "No", "-", "Overall rating score (1-5)"),
            ("quality_rating", "PositiveInt", "No", "-", "Quality score (1-5)"),
            ("communication_rating", "PositiveInt", "No", "-", "Communication score (1-5)"),
            ("delivery_rating", "PositiveInt", "No", "-", "Delivery score (1-5)"),
            ("packaging_rating", "PositiveInt", "No", "-", "Packaging score (1-5)"),
            ("comment", "TextField", "Yes", "-", "Detailed evaluation feedback"),
            ("created_at", "DateTimeField", "No", "-", "Rating timestamp")
        ]),
        ("Table 16: marketplace_feedback (Platform Feedback)", [
            ("id", "BigAutoField", "No", "Primary Key", "Feedback ID"),
            ("user_id", "ForeignKey", "No", "FK (User)", "Submitting user reference"),
            ("subject", "CharField", "No", "-", "Feedback subject line"),
            ("message", "TextField", "No", "-", "Detailed feedback message"),
            ("is_resolved", "BooleanField", "No", "-", "Admin resolution status flag"),
            ("created_at", "DateTimeField", "No", "-", "Submission timestamp")
        ]),
        ("Table 17: marketplace_pricetrend (Crop Market Price Trend)", [
            ("id", "BigAutoField", "No", "Primary Key", "Trend Record ID"),
            ("crop_name", "CharField", "No", "-", "Crop commodity name"),
            ("category_id", "ForeignKey", "No", "FK (Category)", "Category reference"),
            ("avg_price", "DecimalField", "No", "-", "Average market price per kg (INR)"),
            ("record_date", "DateField", "No", "-", "Recording date")
        ]),
        ("Table 18: marketplace_marketinsight (Market Forecast & Analytics)", [
            ("id", "BigAutoField", "No", "Primary Key", "Insight Record ID"),
            ("crop_name", "CharField", "No", "-", "Commodity crop name"),
            ("category_id", "ForeignKey", "No", "FK (Category)", "Category reference"),
            ("current_price", "DecimalField", "No", "-", "Current APMC mandi price (INR)"),
            ("price_change_percent", "DecimalField", "No", "-", "Price percentage change (+/- %)"),
            ("demand_level", "CharField", "No", "-", "Demand index (High/Medium/Low)"),
            ("forecast_direction", "CharField", "No", "-", "Forecast trend direction"),
            ("region", "CharField", "No", "-", "APMC Mandi / Region name"),
            ("volume_tonnes", "IntegerField", "No", "-", "Traded volume in tonnes"),
            ("record_date", "DateField", "No", "-", "Generation record date")
        ]),
        ("Table 19: marketplace_returnrequest (Order Return & Refund Request)", [
            ("id", "BigAutoField", "No", "Primary Key", "Return Request ID"),
            ("order_id", "ForeignKey", "No", "FK (Order)", "Parent order reference"),
            ("reason", "CharField", "No", "-", "Return reason choice"),
            ("description", "TextField", "Yes", "-", "Detailed return reason description"),
            ("status", "CharField", "No", "-", "Status (Pending/Approved/Rejected)"),
            ("created_at", "DateTimeField", "No", "-", "Request filing timestamp")
        ]),
        ("Table 20: marketplace_report (System & Sales Reports)", [
            ("id", "BigAutoField", "No", "Primary Key", "Report Record ID"),
            ("name", "CharField", "No", "-", "Report document title"),
            ("report_type", "CharField", "No", "-", "Report classification type"),
            ("generated_by_id", "ForeignKey", "No", "FK (User)", "Generating user reference"),
            ("format", "CharField", "No", "-", "File format (pdf, csv, xlsx)"),
            ("download_count", "PositiveInt", "No", "-", "Download counter"),
            ("file", "FileField", "Yes", "-", "Generated file path"),
            ("scheduled_interval", "CharField", "Yes", "-", "Scheduled frequency (daily, weekly, monthly)"),
            ("created_at", "DateTimeField", "No", "-", "Report generation timestamp")
        ]),
        ("Table 21: marketplace_orderotp (Order Confirmation OTP)", [
            ("id", "BigAutoField", "No", "Primary Key", "Order OTP Record ID"),
            ("user_id", "ForeignKey", "No", "FK (User)", "Buyer user reference"),
            ("otp_code", "CharField", "No", "-", "6-digit OTP passcode"),
            ("shipping_address", "TextField", "No", "-", "Shipping destination address"),
            ("payment_method", "CharField", "No", "-", "Selected payment method"),
            ("created_at", "DateTimeField", "No", "-", "OTP generation timestamp"),
            ("expires_at", "DateTimeField", "No", "-", "OTP expiration timestamp"),
            ("resend_count", "PositiveInt", "No", "-", "Resend attempt counter"),
            ("is_verified", "BooleanField", "No", "-", "Verification flag status")
        ])
    ]

    for title, fields in models_info:
        add_heading_2(title)
        t = doc.add_table(rows=1, cols=5)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t)
        
        hdr = t.rows[0].cells
        hdr[0].text = "Field Name"
        hdr[1].text = "Data Type"
        hdr[2].text = "Null?"
        hdr[3].text = "Constraint"
        hdr[4].text = "Description"
        set_cell_background(hdr[0], "1B4332")
        set_cell_background(hdr[1], "1B4332")
        set_cell_background(hdr[2], "1B4332")
        set_cell_background(hdr[3], "1B4332")
        set_cell_background(hdr[4], "1B4332")
        for cell in hdr:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

        for fn, dt, nl, cn, ds in fields:
            row_cells = t.add_row().cells
            row_cells[0].text = fn
            row_cells[1].text = dt
            row_cells[2].text = nl
            row_cells[3].text = cn
            row_cells[4].text = ds
            set_cell_background(row_cells[0], "F4F9F4")
            row_cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 9: SCREENSHOTS & MODULE DESCRIPTIONS
    # -------------------------------------------------------------
    add_heading_1("9. Screen Shots & Module Descriptions")
    add_paragraph("This section presents visual screen captures and detailed functional descriptions for all major modules of the AgriConnect system across Farmer, Buyer, and Admin portals.")

    screens_docs = [
        ("screen_login.png", "9.1 Login & Authentication Screen", "Multi-role secure authentication portal supporting login for Farmers, Buyers, and System Administrators with input validation and session tracking."),
        ("screen_admin_dashboard.png", "9.2 Admin Control Dashboard", "Centralized administrative management dashboard displaying platform user metrics, crop approval queues, revenue summary, and system report generation."),
        ("screen_farmer_dashboard.png", "9.3 Farmer Management Dashboard", "Farmer management portal providing harvest stock oversight, incoming buyer orders tracking, price updating tools, and revenue analytics."),
        ("screen_marketplace.png", "9.4 Buyer Marketplace & Crop Catalog", "Dynamic produce catalog presenting available crops with high-resolution imagery, category filtering, search, price per kg, and stock availability."),
        ("screen_crop_detail.png", "9.5 Crop Detail & Quantity Selector", "Detailed crop view displaying harvest date, shelf-life indicators, farmer verification badge, dynamic quantity calculator, and Add-to-Cart controls."),
        ("screen_cart_checkout.png", "9.6 Shopping Cart & Checkout Page", "Shopping cart summary with itemized pricing, delivery address selection, multi-mode payment options (COD/UPI), and OTP verification step."),
        ("screen_farmer_store.png", "9.7 Public Farmer Storefront", "Dedicated farmer profile page highlighting farm location, verification badge, overall rating breakdown, published crop listings, and direct chat option."),
        ("screen_order_history.png", "9.8 Order Tracking & History Page", "Order tracking dashboard allowing buyers and farmers to monitor real-time shipment status, download PDF tax invoices, and submit return requests."),
        ("screen_price_trends.png", "9.9 Market Price Trends & Analytics", "Analytical insight dashboard providing historical crop price trends, market demand forecasts, and price fluctuation statistics to guide farmers."),
        ("screen_chat.png", "9.10 Direct Messaging & Buyer-Farmer Chat", "Real-time communication module enabling buyers and farmers to negotiate prices, clarify delivery details, and exchange product inquiries directly.")
    ]

    for img_file, screen_title, desc in screens_docs:
        add_heading_2(screen_title)
        img_path = f'report_assets/{img_file}'
        if os.path.exists(img_path):
            doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_picture(img_path, width=Inches(5.8))
        add_paragraph(desc)
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 10: TEST CASES
    # -------------------------------------------------------------
    add_heading_1("10. Test Cases")
    add_paragraph("Comprehensive software testing was conducted across all system modules to verify user input validation, data integrity, security controls, and functional correctness. All test cases yielded expected outputs and passed successfully.")

    test_suites = [
        ("Test Suite 1: Login & User Authentication", [
            ("Username & Password", "Display input fields for username & password", "System renders two clean input text fields", "Fields rendered cleanly", "Pass"),
            ("Login Button", "Verify login with valid credentials", "User authenticated & redirected to role dashboard", "Redirected to dashboard", "Pass"),
            ("Validation", "Check empty field submission", "Displays error: 'Username/Password required'", "Error displayed", "Pass"),
            ("Invalid Credentials", "Check login with incorrect password", "Displays error: 'Invalid username or password'", "Error displayed", "Pass")
        ]),
        ("Test Suite 2: Farmer Crop Management", [
            ("Add Crop", "Test adding new crop with price & stock", "New crop added and set to pending approval", "Crop added successfully", "Pass"),
            ("Edit Crop", "Update crop price per kg and stock quantity", "Crop listing updated in database and marketplace", "Listing updated", "Pass"),
            ("Stock Status", "Set crop quantity to 0", "Listing automatically marked 'Out of Stock'", "Marked out of stock", "Pass"),
            ("Delete Crop", "Remove crop listing from farmer portal", "Crop entry removed from catalog", "Removed successfully", "Pass")
        ]),
        ("Test Suite 3: Shopping Cart & Checkout", [
            ("Add to Cart", "Select crop quantity and click Add to Cart", "Crop added to user cart with total calculated", "Added to cart", "Pass"),
            ("Update Quantity", "Increase item quantity in shopping cart", "Subtotal and grand total dynamically update", "Totals updated", "Pass"),
            ("Remove Item", "Click delete icon on cart item", "Item removed from cart", "Item removed", "Pass"),
            ("Checkout OTP", "Submit order with phone number", "System generates 6-digit OTP for order confirmation", "OTP sent & verified", "Pass")
        ]),
        ("Test Suite 4: Admin Control & Approval", [
            ("Crop Approval", "Admin approves pending farmer crop listing", "Crop status changed to Approved & visible in marketplace", "Status updated to Approved", "Pass"),
            ("User Role Audit", "Admin inspects registered users list", "Displays all active farmers, buyers, and admin accounts", "List displayed cleanly", "Pass"),
            ("System Reports", "Click Generate Sales Report PDF", "PDF invoice/report generated and downloaded", "Report generated", "Pass")
        ])
    ]

    for suite_title, tests in test_suites:
        add_heading_2(suite_title)
        t = doc.add_table(rows=1, cols=5)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t)
        
        hdr = t.rows[0].cells
        hdr[0].text = "Field / Module"
        hdr[1].text = "Test Case Description"
        hdr[2].text = "Expected Output"
        hdr[3].text = "Actual Output"
        hdr[4].text = "Result"
        set_cell_background(hdr[0], "1B4332")
        set_cell_background(hdr[1], "1B4332")
        set_cell_background(hdr[2], "1B4332")
        set_cell_background(hdr[3], "1B4332")
        set_cell_background(hdr[4], "1B4332")
        for cell in hdr:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

        for fld, tc_desc, exp, act, res in tests:
            row_cells = t.add_row().cells
            row_cells[0].text = fld
            row_cells[1].text = tc_desc
            row_cells[2].text = exp
            row_cells[3].text = act
            row_cells[4].text = res
            set_cell_background(row_cells[0], "F4F9F4")
            row_cells[0].paragraphs[0].runs[0].font.bold = True
            row_cells[4].paragraphs[0].runs[0].font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 11: LIMITATIONS
    # -------------------------------------------------------------
    add_heading_1("11. Limitations")
    add_bullet("Internet Connectivity Dependency: The application requires stable internet access for real-time market price updates and order placements, which may pose challenges in remote rural regions with limited network coverage.")
    add_bullet("Manual Inventory Updates: Farmers must manually update crop availability and harvested quantities, leading to potential discrepancies if stock updates are delayed after local offline sales.")
    add_bullet("Scalability for Enterprise Volumes: While the current Django architecture easily handles thousands of concurrent users, massive enterprise scaling across multiple geographical zones requires additional cloud load balancing and redis caching.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 12: FUTURE ENHANCEMENTS
    # -------------------------------------------------------------
    add_heading_1("12. Future Enhancements")
    add_bullet("AI-Driven Crop Disease Identification: Integrate machine learning models allowing farmers to upload leaf photos to automatically diagnose crop diseases and receive treatment remedies.")
    add_bullet("Multilingual Regional Voice Support: Introduce voice-assisted navigation in Hindi, Gujarati, Marathi, and regional languages to assist non-literate farmers.")
    add_bullet("Weather & Price Predictive Analytics: Incorporate real-time weather API feeds and AI price prediction forecasting to help farmers choose optimal harvest times.")
    add_bullet("Native Mobile Application: Expand AgriConnect to native Android and iOS mobile applications using Flutter for seamless mobile notifications and offline capability.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 13: WEBLIOGRAPHY
    # -------------------------------------------------------------
    add_heading_1("13. Webliography")
    add_bullet("Django Web Framework Documentation: https://docs.djangoproject.com/")
    add_bullet("Python Official Documentation: https://www.python.org/")
    add_bullet("Bootstrap 5 Documentation: https://getbootstrap.com/")
    add_bullet("Stack Overflow Technical Q&A: https://stackoverflow.com/")
    add_bullet("Draw.io Diagramming Tool: https://app.diagrams.net/")
    add_bullet("MDN Web Docs (HTML/CSS/JS): https://developer.mozilla.org/")

    doc.save("AgriConnect_Project_Report.docx")
    doc.save("AgriConnect_Project_Report_v2.docx")
    print("AgriConnect_Project_Report.docx & AgriConnect_Project_Report_v2.docx generated successfully!")

if __name__ == "__main__":
    create_report()
